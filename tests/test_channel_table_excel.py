"""A channel table is a spreadsheet, and now goes both ways.

Writing one was never the hard part. Reading one back is what makes the
export worth having: a channel table arrives as a spreadsheet far more
often than as anything else — a calibration lab sends one, a controller
writes one — and until this existed, a table exported from the window
could not be brought back into it.
"""

from __future__ import annotations

import numpy as np
import pytest
from openpyxl import Workbook

import visualdynamics
from visualdynamics.core.channel_table import ChannelTable
from visualdynamics.io import excel


@pytest.fixture
def table():
    return ChannelTable({
        'channel': [1, 2],
        'node': ['101', '102'],
        'direction': ['Z+', 'X+'],
        'unit': ['g', 'g'],
        # the two that a numeric read would ruin
        'serial_number': ['007', 'LW-4432'],
        'sensitivity': ['10.02', '9.98'],
    })


def test_a_table_survives_the_round_trip_exactly(table, tmp_path):
    path = str(tmp_path / 'channels.xlsx')
    visualdynamics.export_file(table, path, format='excel')
    back = visualdynamics.import_file(path)
    assert isinstance(back, ChannelTable)
    assert back == table, 'the object itself, not merely something like it'
    assert list(back['serial_number']) == ['007', 'LW-4432'], (
        "'007' read as a number comes back '7'"
    )


def written(tmp_path, rows, name='book.xlsx'):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    path = str(tmp_path / name)
    book.save(path)
    return path


def test_the_header_need_not_be_the_first_row(tmp_path):
    """A controller writes merged group headings above the column names.

    Rather than teach this two layouts, it looks for the row that *names
    columns* — a spreadsheet's header is the first row saying 'channel'.
    """
    path = written(tmp_path, [
        ['Physical Device', None, None, 'Calibration'],
        ['Channel', 'Node Number', 'Direction', 'Serial Number'],
        [1, '101', 'Z+', '007'],
        [2, '102', 'X+', 'LW-4432'],
    ])
    back = visualdynamics.import_file(path)
    assert back.num_channels == 2
    assert list(back['serial_number']) == ['007', 'LW-4432']


def test_other_spellings_of_the_four_core_columns_land_in_them(tmp_path):
    """'Node Number' means node. Left alone it arrives as an *extra*
    column while `node` is created empty beside it — the same table with
    its DOFs thrown away."""
    path = written(tmp_path, [
        ['Ch', 'Point', 'Axis', 'Units'],
        [1, '101', 'Z+', 'g'],
    ])
    back = visualdynamics.import_file(path)
    # the schema leads in its own order; the spellings must have landed
    # in the schema's columns, not created extras beside them
    assert back.column_names[:len(back.SCHEMA)] == list(back.SCHEMA)
    assert back.dof_strings() == ['101Z+']
    assert list(back['unit']) == ['g']


def test_a_known_column_lands_and_an_invented_one_does_not(tmp_path):
    """'Cal Due' is the expiration under another name and is kept; a
    rack slot is the lab's own bookkeeping and is not a property of the
    measurement, so the fixed schema lets it go."""
    path = written(tmp_path, [
        ['Channel', 'Cal Due', 'Rack Slot'],
        [1, '2027-01-04', 'A3'],
    ])
    back = visualdynamics.import_file(path)
    assert list(back['expiration']) == ['2027-01-04']
    assert 'rack_slot' not in back.column_names
    assert back.column_names == list(back.SCHEMA)


def test_blank_rows_below_the_table_are_not_channels(tmp_path):
    path = written(tmp_path, [
        ['Channel', 'Node'], [1, '101'], [None, None], [None, None],
    ])
    assert visualdynamics.import_file(path).num_channels == 1


def test_a_table_with_no_dofs_yet_still_imports(tmp_path):
    """Somebody forgot to record the node. Importing it is how they get
    to fix it; refusing is how it stays broken in a spreadsheet."""
    path = written(tmp_path, [['Channel', 'Unit'], [1, 'g'], [2, 'g']])
    back = visualdynamics.import_file(path)
    assert back.num_channels == 2
    assert list(back['node']) == ['', '']


def test_a_spreadsheet_that_is_not_a_channel_table_is_left_alone(tmp_path):
    """This importer is asked about every file dropped on the window, so
    a wrong yes is expensive."""
    path = written(tmp_path, [['Date', 'Amount'], ['2026-01-01', 12.5]])
    assert not excel.sniff(path)
    with pytest.raises(ValueError, match='No importer recognizes'):
        visualdynamics.import_file(path)


def test_something_that_is_not_a_spreadsheet_at_all(tmp_path):
    path = tmp_path / 'not-really.xlsx'
    path.write_bytes(b'this is not a workbook')
    assert not excel.sniff(str(path)), 'sniff must not raise on rubbish'


def test_the_imported_table_carries_into_a_project(tmp_path):
    """The point of importing it: a table read from a lab's spreadsheet
    has to behave like one built in the window."""
    path = written(tmp_path, [
        ['Channel', 'Node', 'Direction', 'Unit'],
        [1, '101', 'Z+', 'g'], [2, '101', 'X+', 'g'],
    ])
    project = visualdynamics.Project('T')
    project.add('Channels', visualdynamics.import_file(path))
    assert project.channel_table.dof_strings() == ['101Z+', '101X+']
    saved = str(tmp_path / 'p.vdyn')
    project.save(saved)
    assert visualdynamics.Project.open(saved)['Channels'].num_channels == 2
    assert np.all(
        visualdynamics.Project.open(saved)['Channels']['unit'] == 'g')


def test_the_export_carries_excels_own_drop_downs(tmp_path):
    """A channel table is very often filled in by somebody else — a
    calibration lab, an engineer with the spreadsheet and no copy of
    this — so the constraints have to travel with the file. Excel
    refuses a bad direction there for the same reason `set_cell`
    refuses one here."""
    from openpyxl import load_workbook

    from visualdynamics.core.channel_table import ChannelTable

    table = ChannelTable({'channel': [1, 2], 'node': ['101', '102'],
                          'direction': ['Z+', 'X+'], 'unit': ['g', 'g']})
    path = str(tmp_path / 'channels.xlsx')
    visualdynamics.export_file(table, path, format='excel')
    sheet = load_workbook(path).active
    lists = {str(dv.sqref): dv.formula1
             for dv in sheet.data_validations.dataValidation}
    assert lists, 'no drop-downs were written'
    # every choice column got one, over the data rows only
    every = ' '.join(lists.values())
    for choice in ('X+', 'RZ-', 'reference', 'monitor', 'True', 'False'):
        assert choice in every, choice
    assert all(':' in ref for ref in lists), (
        'a validation covers the data rows, not one cell')
    # the sheet says the word the interface says
    assert 'displacement' in every and '"length' not in every


def test_a_sheet_saying_displacement_reads_back_as_length(tmp_path):
    """The export writes the display word, so the import has to take
    it: a spreadsheet reading 'length' beside an app reading
    'displacement' is the kind of difference nobody forgives."""
    path = written(tmp_path, [
        ['Channel', 'Node', 'Direction', 'Channel Type', 'Unit'],
        [1, '101', 'Z+', 'displacement', 'mm'],
    ])
    back = visualdynamics.import_file(path)
    assert list(back['channel_type']) == ['length']
    assert back.units_for(0) == ['mil', 'in', 'm', 'mm']


def test_numbers_are_written_as_numbers(tmp_path):
    """Excel puts a green triangle and a 'number stored as text'
    warning on every cell otherwise, and a node is always a number."""
    from openpyxl import load_workbook

    from visualdynamics.core.channel_table import ChannelTable

    table = ChannelTable({'channel': [1, 2], 'node': ['101', ''],
                          'direction': ['Z+', 'X+'], 'unit': ['g', 'g']})
    path = str(tmp_path / 'channels.xlsx')
    visualdynamics.export_file(table, path, format='excel')
    sheet = load_workbook(path)['Channel Table']
    names = table.column_names
    channel_at, node_at = names.index('channel') + 1, names.index('node') + 1
    assert isinstance(sheet.cell(2, channel_at).value, int)
    assert isinstance(sheet.cell(2, node_at).value, int)
    assert sheet.cell(3, node_at).value is None, 'a blank node stays blank'
    assert list(visualdynamics.import_file(path)['node']) == ['101', '']


def test_the_unit_drop_down_follows_the_channel_type(tmp_path):
    """Excel cannot hold a different inline list per row, but it can
    look one up: the units for each quantity go on a hidden sheet as
    named ranges, and the rule is INDIRECT of the type cell beside it."""
    from openpyxl import load_workbook

    from visualdynamics.core.channel_table import ChannelTable
    from visualdynamics.core.unit_choices import ORDINATE_UNITS

    table = ChannelTable({'channel': [1], 'node': ['101'],
                          'direction': ['Z+'], 'unit': ['V'],
                          'channel_type': ['voltage']})
    path = str(tmp_path / 'channels.xlsx')
    visualdynamics.export_file(table, path, format='excel')
    book = load_workbook(path)
    assert book['Units'].sheet_state == 'hidden'
    # one named range per quantity, plus the blank-type fallback
    assert 'voltage' in book.defined_names
    assert 'AllUnits' in book.defined_names
    sheet = book['Channel Table']
    unit_at = table.column_names.index('unit') + 1
    rules = [dv for dv in sheet.data_validations.dataValidation
             if str(dv.sqref).startswith(
                 sheet.cell(1, unit_at).column_letter)]
    assert len(rules) == 1
    assert rules[0].formula1.startswith('=INDIRECT('), rules[0].formula1
    # and the hidden sheet holds that quantity's units to point at
    written = {cell.value for col in book['Units'].columns
               for cell in col if cell.value}
    assert set(ORDINATE_UNITS['voltage']) <= written


def test_strain_has_units_of_its_own(tmp_path):
    """It is a quantity a channel measures, and offering a strain gauge
    every unit in the list was what the missing entry did."""
    from visualdynamics.core.channel_table import ChannelTable
    from visualdynamics.core.unit_choices import ORDINATE_UNITS

    table = ChannelTable({'channel': [1], 'node': ['101'],
                          'direction': ['Z+'], 'unit': [''],
                          'channel_type': ['strain']})
    assert table.units_for(0) == ORDINATE_UNITS['strain']
    assert 'm/s**2' not in table.units_for(0)
    table.set_cell('unit', 0, 'microstrain')
